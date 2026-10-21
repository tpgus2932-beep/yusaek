from __future__ import annotations

import io
from pathlib import Path

import openpyxl

from services.order_recommendation_store import now_kst_iso

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "assets" / "luna_order_template.xlsx"
HEADER_ROW = 2
FIRST_DATA_ROW = 3
CODE_COL, NAME_COL, OPTION_COL, QTY_COL = 2, 3, 4, 5


def _norm(value) -> str:
    return "" if value is None else str(value).strip()


def _init_table(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS luna_order_mapping (
            yusas_code TEXT PRIMARY KEY,
            template_code TEXT NOT NULL,
            template_name TEXT NOT NULL,
            template_option TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            updated_by TEXT NOT NULL DEFAULT ''
        )
        """
    )
    conn.commit()


def load_template_rows() -> list[dict]:
    """양식의 상품 행(품목코드/상품명(C열)/옵션명(D열))을 읽는다. 열 순서가 고정이라 인덱스로 읽는다."""
    wb = openpyxl.load_workbook(TEMPLATE_PATH, read_only=True)
    try:
        ws = wb.worksheets[0]
        rows = []
        for idx, row in enumerate(ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), start=FIRST_DATA_ROW):
            name = _norm(row[NAME_COL - 1]) if len(row) >= NAME_COL else ""
            if not name:
                continue
            rows.append({
                "row": idx,
                "code": _norm(row[CODE_COL - 1]),
                "name": name,
                "option": _norm(row[OPTION_COL - 1]) if len(row) >= OPTION_COL else "",
            })
        return rows
    finally:
        wb.close()


def _template_key(code, name, option) -> tuple[str, str, str]:
    return (_norm(code), _norm(name), _norm(option))


def list_mappings(get_shared_db) -> list[dict]:
    conn = get_shared_db()
    try:
        _init_table(conn)
        rows = conn.execute(
            "SELECT yusas_code, template_code, template_name, template_option, updated_at, updated_by "
            "FROM luna_order_mapping ORDER BY updated_at DESC"
        ).fetchall()
        return [{k: r[k] for k in r.keys()} for r in rows]
    finally:
        conn.close()


def save_mapping(get_shared_db, *, yusas_code: str, template_code: str, template_name: str,
                 template_option: str, user: str) -> None:
    conn = get_shared_db()
    try:
        _init_table(conn)
        conn.execute(
            """
            INSERT INTO luna_order_mapping
                (yusas_code, template_code, template_name, template_option, updated_at, updated_by)
            VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(yusas_code) DO UPDATE SET
                template_code = excluded.template_code,
                template_name = excluded.template_name,
                template_option = excluded.template_option,
                updated_at = excluded.updated_at,
                updated_by = excluded.updated_by
            """,
            (yusas_code, _norm(template_code), _norm(template_name), _norm(template_option), now_kst_iso(), user),
        )
        conn.commit()
    finally:
        conn.close()


def delete_mapping(get_shared_db, yusas_code: str) -> None:
    conn = get_shared_db()
    try:
        _init_table(conn)
        conn.execute("DELETE FROM luna_order_mapping WHERE yusas_code = ?", (yusas_code,))
        conn.commit()
    finally:
        conn.close()


def resolve_items(get_shared_db, items: list[dict]) -> dict:
    """발주 항목을 저장된 매칭으로 양식 행에 연결한다. 자동 추측은 하지 않는다 -
    한 번이라도 수동으로 잡아둔 yusas_code만 매칭되고, 나머지는 unmatched로 돌려준다.
    매칭은 저장돼 있어도 양식에서 그 행이 사라졌으면(양식 개정 등) unmatched다."""
    template_rows = load_template_rows()
    row_by_key = {_template_key(r["code"], r["name"], r["option"]): r["row"] for r in template_rows}
    mapping_by_code = {m["yusas_code"]: m for m in list_mappings(get_shared_db)}

    matched, unmatched = [], []
    for item in items:
        mapping = mapping_by_code.get(item.get("yusas_code"))
        row = None
        if mapping:
            row = row_by_key.get(_template_key(mapping["template_code"], mapping["template_name"], mapping["template_option"]))
        if row is None:
            unmatched.append(item)
        else:
            matched.append({**item, "row": row})
    return {"matched": matched, "unmatched": unmatched}


def build_workbook(matched: list[dict]) -> bytes:
    """양식 E열(주문수량)에 매칭된 항목 수량을 합산해서 채운 엑셀을 만든다.
    같은 행에 여러 항목이 매칭되거나 E열에 이미 값이 있으면 덮어쓰지 않고 더한다."""
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.worksheets[0]
    for item in matched:
        qty = int(item.get("qty") or 0)
        if qty <= 0:
            continue
        cell = ws.cell(row=item["row"], column=QTY_COL)
        existing = cell.value if isinstance(cell.value, (int, float)) else 0
        cell.value = existing + qty
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
